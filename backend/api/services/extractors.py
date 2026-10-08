import base64
import csv
import io
import json
import logging
import re
from pathlib import PurePosixPath
from urllib.parse import unquote, urlsplit

from bs4 import BeautifulSoup
from django.conf import settings

from ..exceptions import ScrapeError

logger = logging.getLogger(__name__)

HTML_TYPES = ("text/html", "application/xhtml+xml")
TEXT_TYPES = ("text/plain", "text/markdown", "text/x-markdown")
CSV_TYPES = ("text/csv", "application/csv")
JSON_TYPES = ("application/json", "application/ld+json", "text/json")
XML_TYPES = ("application/xml", "text/xml", "application/rss+xml", "application/atom+xml")
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
IMAGE_TYPES = ("image/png", "image/jpeg", "image/webp", "image/gif", "image/heic", "image/heif")

EXTENSIONS = {
    ".html": "html",
    ".htm": "html",
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "text",
    ".md": "text",
    ".markdown": "text",
    ".csv": "csv",
    ".json": "json",
    ".xml": "xml",
    ".rss": "xml",
    ".atom": "xml",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".webp": "image",
    ".gif": "image",
    ".mp4": "media",
    ".webm": "media",
    ".mov": "media",
    ".mp3": "media",
    ".wav": "media",
    ".m4a": "media",
}

READABLE_WORD = re.compile(r"[^\W\d_]{2,}", re.UNICODE)


def clean(text):
    return " ".join(text.split())


def file_name(url):
    name = PurePosixPath(unquote(urlsplit(url).path)).name
    return name or urlsplit(url).hostname or url


def detect_kind(content_type, url, body):
    mime = (content_type or "").split(";")[0].strip().lower()
    if body.startswith(b"%PDF-") or mime == "application/pdf":
        return "pdf"
    if mime == DOCX_TYPE:
        return "docx"
    if mime in IMAGE_TYPES:
        return "image"
    if mime.startswith(("video/", "audio/")):
        return "media"
    if mime in HTML_TYPES:
        return "html"
    if mime in JSON_TYPES or mime.endswith("+json"):
        return "json"
    if mime in XML_TYPES or mime.endswith("+xml"):
        return "xml"
    if mime in CSV_TYPES:
        return "csv"
    if mime in TEXT_TYPES:
        return "text"

    by_extension = EXTENSIONS.get(PurePosixPath(urlsplit(url).path).suffix.lower())
    if by_extension:
        return by_extension
    head = body[:512].lstrip().lower()
    if head.startswith((b"<!doctype html", b"<html")) or b"<body" in head:
        return "html"
    if mime.startswith("text/") or not mime:
        return "text"
    return "unknown"


def decode(body, charset=None):
    for encoding in filter(None, (charset, "utf-8", "cp1252")):
        try:
            return body.decode(encoding)
        except (LookupError, UnicodeDecodeError):
            continue
    return body.decode("utf-8", errors="replace")


def page_title(soup):
    og_title = soup.find("meta", property="og:title")
    if og_title and og_title.get("content"):
        return clean(og_title["content"])
    if soup.title:
        return clean(soup.title.get_text())
    return ""


def html_text(html, encoding=None):
    soup = BeautifulSoup(html, settings.SCRAPER_HTML_PARSER, from_encoding=encoding if isinstance(html, bytes) else None)
    title = page_title(soup)

    for tag in soup(settings.SCRAPER_IGNORED_TAGS):
        tag.decompose()

    content = soup.find("article") or soup.find("main") or soup.body or soup

    lines = []
    seen = set()
    for el in content.find_all(settings.SCRAPER_TEXT_TAGS):
        text = clean(el.get_text(" "))
        if len(text) > 1 and text not in seen:
            seen.add(text)
            lines.append(text)
    text = "\n".join(lines)

    if len(text) < settings.SCRAPER_FALLBACK_MIN_CHARS:
        logger.debug("Not many text tags found, falling back to all text")
        text = clean(content.get_text(" "))

    return title[: settings.SCRAPER_MAX_TITLE_CHARS], text


def looks_like_prose(value):
    if len(value) < 40 or value.startswith(("http://", "https://", "/", "data:")):
        return False
    if any(mark in value for mark in ("{", "}", "</", "function(", "=>")):
        return False
    words = READABLE_WORD.findall(value)
    return len(words) >= 6 and sum(len(w) for w in words) > len(value) * 0.5


def prose_strings(data, found, limit=400):
    if len(found) >= limit:
        return
    if isinstance(data, str):
        text = clean(BeautifulSoup(data, "html.parser").get_text(" ")) if "<" in data else clean(data)
        if looks_like_prose(text) and text not in found:
            found.append(text)
    elif isinstance(data, dict):
        for value in data.values():
            prose_strings(value, found, limit)
    elif isinstance(data, list):
        for value in data:
            prose_strings(value, found, limit)


def embedded_text(html, encoding=None):
    soup = BeautifulSoup(html, settings.SCRAPER_HTML_PARSER, from_encoding=encoding if isinstance(html, bytes) else None)
    found = []

    for name in ("description", "og:description", "twitter:description"):
        meta = soup.find("meta", attrs={"name": name}) or soup.find("meta", property=name)
        if meta and meta.get("content"):
            text = clean(meta["content"])
            if text and text not in found:
                found.append(text)

    for script in soup.find_all("script"):
        kind = (script.get("type") or "").lower()
        if kind not in ("application/ld+json", "application/json") and script.get("id") not in ("__NEXT_DATA__",):
            continue
        try:
            prose_strings(json.loads(script.string or ""), found)
        except ValueError:
            continue

    noscript = clean(" ".join(n.get_text(" ") for n in soup.find_all("noscript")))
    if looks_like_prose(noscript):
        found.append(noscript)

    return page_title(soup)[: settings.SCRAPER_MAX_TITLE_CHARS], "\n".join(found)


def pdf_text(body):
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(body))
        if reader.is_encrypted:
            reader.decrypt("")
        pages = reader.pages[: settings.SCRAPER_MAX_PDF_PAGES]
        text = "\n".join(clean(page.extract_text() or "") for page in pages)
        title = clean(str((reader.metadata or {}).get("/Title") or ""))
    except (PdfReadError, ValueError, KeyError, TypeError) as e:
        logger.warning("Could not read PDF: %r", e)
        raise ScrapeError("We couldn't read that PDF. It may be damaged or protected.", 422, "bad_file") from None
    if len(reader.pages) > settings.SCRAPER_MAX_PDF_PAGES:
        logger.info("PDF has %d pages, reading the first %d", len(reader.pages), settings.SCRAPER_MAX_PDF_PAGES)
    return title, text


def docx_text(body):
    from docx import Document

    try:
        document = Document(io.BytesIO(body))
    except Exception as e:
        logger.warning("Could not read Word file: %r", e)
        raise ScrapeError("We couldn't read that Word document.", 422, "bad_file") from None
    lines = [clean(p.text) for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(clean(cell.text) for cell in row.cells))
    return clean(document.core_properties.title or ""), "\n".join(lines)


def json_text(body, charset=None):
    try:
        data = json.loads(decode(body, charset))
    except ValueError:
        return "", decode(body, charset)
    return "", json.dumps(data, ensure_ascii=False, indent=1)


def csv_text(body, charset=None):
    rows = list(csv.reader(io.StringIO(decode(body, charset))))
    return "", "\n".join(" | ".join(cell.strip() for cell in row) for row in rows if any(row))


def xml_text(body):
    soup = BeautifulSoup(body, "xml")
    title = soup.find("title")
    lines = []
    for item in soup.find_all(["item", "entry"]) or [soup]:
        parts = [clean(tag.get_text(" ")) for tag in item.find_all(["title", "description", "summary", "content"])]
        line = " - ".join(p for p in parts if p) or clean(item.get_text(" "))
        if line:
            lines.append(clean(BeautifulSoup(line, "html.parser").get_text(" ")))
    return clean(title.get_text()) if title else "", "\n".join(lines)


def image_data(body, content_type):
    if len(body) > settings.SCRAPER_MAX_IMAGE_BYTES:
        limit = settings.SCRAPER_MAX_IMAGE_BYTES // (1024 * 1024)
        raise ScrapeError(f"That image is too large. The limit is {limit} MB.", 413, "file_too_large")
    mime = (content_type or "").split(";")[0].strip().lower()
    if mime not in IMAGE_TYPES:
        mime = "image/png" if body.startswith(b"\x89PNG") else "image/jpeg"
    return {"mime_type": mime, "data": base64.b64encode(body).decode()}
